# /home/shlicha/psycho_diary/bot/services/export_service.py
"""
Сервис для экспорта данных пользователя в формате .docx
"""

import os
import tempfile
from datetime import datetime, timedelta
from asgiref.sync import sync_to_async

from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from django.utils import timezone

from users.models import User
from diary.models import DiaryEntry, TherapyTask, TaskCompletion


async def generate_user_docx(user: User) -> str:
    """
    Генерирует .docx файл с данными пользователя.
    Возвращает путь к файлу.
    """
    # Создаём временный файл
    temp_file = tempfile.NamedTemporaryFile(delete=False, suffix='.docx')
    temp_path = temp_file.name
    temp_file.close()

    doc = Document()

    # --- Заголовок ---
    title = doc.add_heading('ДАННЫЕ ПОЛЬЗОВАТЕЛЯ', 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Дата экспорта
    export_date = datetime.now().strftime('%d.%m.%Y')
    doc.add_paragraph(f'Дата экспорта: {export_date}')
    doc.add_paragraph()

    # --- Психологический профиль ---
    doc.add_heading('ПСИХОЛОГИЧЕСКИЙ ПРОФИЛЬ', level=1)

    # Описание
    if user.personality_description:
        doc.add_heading('Описание личности', level=2)
        doc.add_paragraph(user.personality_description)
        doc.add_paragraph()

    # Черты характера
    if user.personality_traits:
        doc.add_heading('Черты характера', level=2)
        for trait, value in user.personality_traits.items():
            doc.add_paragraph(f'• {trait}: {value}')
        doc.add_paragraph()

    # Факты о жизни
    if user.life_facts:
        doc.add_heading('Факты о жизни', level=2)
        for fact in user.life_facts:
            doc.add_paragraph(f'• {fact}')
        doc.add_paragraph()

    # Увлечения
    if user.hobbies:
        doc.add_heading('Увлечения', level=2)
        for hobby in user.hobbies:
            doc.add_paragraph(f'• {hobby}')
        doc.add_paragraph()

    # Ценности
    if user.core_values:
        doc.add_heading('Ключевые ценности', level=2)
        for value in user.core_values:
            doc.add_paragraph(f'• {value}')
        doc.add_paragraph()

    # Триггеры
    if user.triggers:
        doc.add_heading('Триггеры', level=2)
        for trigger in user.triggers:
            doc.add_paragraph(f'• {trigger}')
        doc.add_paragraph()

    # Сильные стороны
    if user.strengths:
        doc.add_heading('Сильные стороны', level=2)
        for strength in user.strengths:
            doc.add_paragraph(f'• {strength}')
        doc.add_paragraph()

    # Зоны роста
    if user.areas_for_growth:
        doc.add_heading('Зоны роста', level=2)
        for growth in user.areas_for_growth:
            doc.add_paragraph(f'• {growth}')
        doc.add_paragraph()

    # Инсайты
    if user.personality_insights:
        doc.add_heading('Глубинные инсайты', level=2)
        for insight_key, insight_value in user.personality_insights.items():
            doc.add_paragraph(f'• {insight_key}: {insight_value}')
        doc.add_paragraph()

    # Привычки и навыки
    if user.practical_patterns:
        doc.add_heading('Привычки и навыки', level=2)
        for pattern in user.practical_patterns:
            doc.add_paragraph(f'• {pattern}')
        doc.add_paragraph()

    # Стремления
    if user.aspirations:
        doc.add_heading('Стремления', level=2)
        for aspiration in user.aspirations:
            doc.add_paragraph(f'• {aspiration}')
        doc.add_paragraph()

    # --- Записи за последние 30 дней ---
    thirty_days_ago = timezone.now() - timedelta(days=30)

    # Записи
    @sync_to_async
    def get_entries():
        return list(
            DiaryEntry.objects.filter(
                user=user,
                created_at__gte=thirty_days_ago
            ).order_by('created_at')
        )

    entries = await get_entries()

    if entries:
        doc.add_heading('ЗАПИСИ ЗА ПОСЛЕДНИЕ 30 ДНЕЙ', level=1)

        for entry in entries:
            date_str = entry.created_at.strftime('%d.%m.%Y %H:%M')
            doc.add_heading(f'Запись от {date_str}', level=2)

            # Текст записи
            if entry.text:
                doc.add_paragraph(entry.text)
            # Уточняющие вопросы
            if entry.clarification_questions_asked:
                doc.add_heading('Уточняющие вопросы', level=3)
                for q in entry.clarification_questions_asked:
                    doc.add_paragraph(f'• {q}')
            # Ответы на вопросы
            if entry.clarification_answers:
                doc.add_heading('Ответы на уточнения', level=3)
                for q, a in entry.clarification_answers.items():
                    doc.add_paragraph(f'• {q}: {a}')
            # Комментарий пользователя
            if entry.user_comment:
                doc.add_heading('Комментарий пользователя', level=3)
                doc.add_paragraph(entry.user_comment)

            doc.add_paragraph()

    # --- Задания и выполнения ---
    @sync_to_async
    def get_tasks_and_completions():
        tasks = list(
            TherapyTask.objects.filter(
                user=user,
                assigned_at__gte=thirty_days_ago
            ).order_by('assigned_at')
        )
        completions = list(
            TaskCompletion.objects.filter(
                task__user=user,
                created_at__gte=thirty_days_ago
            ).select_related('task')
        )
        return tasks, completions

    tasks, completions = await get_tasks_and_completions()

    if tasks or completions:
        doc.add_heading('ЗАДАНИЯ И ВЫПОЛНЕНИЯ', level=1)

        # Группируем выполнения по заданиям
        completion_map = {}
        for comp in completions:
            task_id = comp.task.id
            if task_id not in completion_map:
                completion_map[task_id] = []
            completion_map[task_id].append(comp)

        for task in tasks:
            date_str = task.assigned_at.strftime('%d.%m.%Y %H:%M')
            doc.add_heading(f'Задание: {task.title}', level=2)
            doc.add_paragraph(f'Дата: {date_str}')
            doc.add_paragraph(f'Описание: {task.description}')
            doc.add_paragraph(f'Статус: {task.get_status_display()}')
            doc.add_paragraph(f'Тип: {task.get_type_display()}')
            doc.add_paragraph(f'Сложность: {task.get_complexity_display()}')

            # Проверяем, есть ли выполнения для этого задания
            task_completions = completion_map.get(task.id, [])
            if task_completions:
                doc.add_heading('Выполнение', level=3)
                for comp in task_completions:
                    comp_date_str = comp.created_at.strftime('%d.%m.%Y %H:%M')
                    doc.add_paragraph(f'Дата выполнения: {comp_date_str}')
                    doc.add_paragraph(f'Текст выполнения: {comp.completion_text}')
                    doc.add_paragraph(f'Ощущения: {comp.feelings_text}')
            else:
                doc.add_paragraph('Выполнение: не выполнено')

            doc.add_paragraph()

    # Сохраняем документ
    doc.save(temp_path)
    return temp_path